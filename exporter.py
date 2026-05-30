from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

# ─── COLORS ───────────────────────────────────────────────
GREEN        = "1E7E34"
BLUE         = "1A56DB"
ORANGE       = "D97706"
GRAY         = "F3F4F6"
WHITE        = "FFFFFF"
DARK         = "111827"
LIGHT_GREEN  = "D1FAE5"
LIGHT_BLUE   = "DBEAFE"
LIGHT_ORANGE = "FEF3C7"

def _thin_border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

def _light_border():
    s = Side(style="thin", color="E5E7EB")
    return Border(left=s, right=s, top=s, bottom=s)

def header_style(cell, color=BLUE):
    cell.font = Font(bold=True, color=WHITE, size=11)
    cell.fill = PatternFill("solid", start_color=color)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = _thin_border()

def data_style(cell, bg=WHITE):
    cell.fill = PatternFill("solid", start_color=bg)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border = _light_border()


def _sheet_summary(wb, result):
    ws = wb.active
    ws.title = "Summary"
    ws.row_dimensions[1].height = 40

    ws["A1"] = "BPO Lead Extraction Report"
    ws["A1"].font = Font(bold=True, size=16, color=DARK)
    ws.merge_cells("A1:B1")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A1"].fill = PatternFill("solid", start_color=GRAY)

    ws["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws["A2"].font = Font(italic=True, color="6B7280", size=10)
    ws.merge_cells("A2:B2")

    rows = [
        ("Company",           result["company"]),
        ("Domain",            result["domain"]),
        ("Website",           result["website"]),
        ("Pages Scraped",     len(result["scraped_pages"])),
        ("Direct Emails",     len(result["direct_emails"])),
        ("Team Members",      len(result["team_members"])),
        ("Verified Emails",   sum(1 for m in result["team_members"] if m.get("verified_email"))),
    ]
    for i, (k, v) in enumerate(rows, start=4):
        ws[f"A{i}"] = k
        ws[f"B{i}"] = str(v)
        ws[f"A{i}"].font = Font(bold=True, color=DARK)
        ws[f"A{i}"].fill = PatternFill("solid", start_color=GRAY)
        data_style(ws[f"B{i}"])

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 45


def _sheet_direct_emails(wb, result):
    ws = wb.create_sheet("Direct Emails")
    ws.row_dimensions[1].height = 32

    for col, h in enumerate(["#", "Email Address", "Company", "Domain"], 1):
        header_style(ws.cell(row=1, column=col, value=h), BLUE)

    for i, email in enumerate(result["direct_emails"], 1):
        bg = LIGHT_BLUE if i % 2 == 0 else WHITE
        for col, val in enumerate([i, email, result["company"], result["domain"]], 1):
            data_style(ws.cell(row=i+1, column=col, value=val), bg)

    for col, w in enumerate([6, 38, 18, 18], 1):
        ws.column_dimensions[get_column_letter(col)].width = w


def _sheet_team(wb, result):
    ws = wb.create_sheet("Team & Emails")
    ws.row_dimensions[1].height = 32

    headers = ["#", "Full Name", "Designation", "Verified Email", "Status", "All Guessed Emails"]
    for col, h in enumerate(headers, 1):
        header_style(ws.cell(row=1, column=col, value=h), GREEN)

    for i, m in enumerate(result["team_members"], 1):
        verified = m.get("verified_email", "")
        status   = m.get("status", "⚠️ Unverified")
        bg = LIGHT_GREEN if verified else (LIGHT_ORANGE if i % 2 == 0 else WHITE)
        values = [i, m["name"], m.get("designation",""), verified, status,
                  ", ".join(m.get("guessed_emails", []))]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=i+1, column=col, value=val)
            data_style(cell, bg)
            if col == 5:
                cell.font = Font(bold=True, color=GREEN if "✅" in str(val) else ORANGE)

    for col, w in enumerate([6, 22, 28, 32, 16, 55], 1):
        ws.column_dimensions[get_column_letter(col)].width = w


def _sheet_outreach(wb, result):
    """সব valid email একসাথে — outreach এর জন্য ready"""
    ws = wb.create_sheet("Outreach List")
    ws.row_dimensions[1].height = 32

    headers = ["#", "Email", "Name", "Designation", "Company", "Domain", "Type", "Status"]
    for col, h in enumerate(headers, 1):
        header_style(ws.cell(row=1, column=col, value=h), ORANGE)

    row_num = 2

    # Direct emails
    for email in result["direct_emails"]:
        bg = LIGHT_BLUE if row_num % 2 == 0 else WHITE
        values = [row_num-1, email, "-", "General/Contact",
                  result["company"], result["domain"], "Direct", "✅ Found on site"]
        for col, val in enumerate(values, 1):
            data_style(ws.cell(row=row_num, column=col, value=val), bg)
        row_num += 1

    # Verified team member emails only
    for m in result["team_members"]:
        if m.get("verified_email"):
            values = [row_num-1, m["verified_email"], m["name"], m.get("designation",""),
                      result["company"], result["domain"], "Team Member", "✅ SMTP Verified"]
            for col, val in enumerate(values, 1):
                data_style(ws.cell(row=row_num, column=col, value=val), LIGHT_GREEN)
            row_num += 1

    for col, w in enumerate([6, 35, 22, 28, 16, 18, 14, 18], 1):
        ws.column_dimensions[get_column_letter(col)].width = w


def save_to_excel(result: dict, output_path: str):
    wb = Workbook()
    _sheet_summary(wb, result)
    _sheet_direct_emails(wb, result)
    _sheet_team(wb, result)
    _sheet_outreach(wb, result)
    wb.save(output_path)
    print(f"\n✅ Excel saved: {output_path}")